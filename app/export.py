"""PDF and DOCX export.

Input is the restricted markdown that app/reports.py produces (and that the
frontend builds for single-incident briefs): `#`/`##`/`###` headings,
`- ` bullets, `1. ` numbered lines, plain paragraphs, and inline **bold**,
*italic* and [text](url) links. Nothing more is supported on purpose - a
tiny parser we fully control is easier to trust (and test) than a general
markdown engine.

Known limitation: the PDF uses a built-in Latin font, so characters outside
Western European text (e.g. Arabic or Chinese headlines from GDELT) print as
"?". DOCX keeps all characters.
"""

import io
import re
from xml.sax.saxutils import escape

_INLINE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|\[(.+?)\]\((.+?)\)")
_NUMBERED = re.compile(r"^(\d+)\.\s+(.*)$")


def parse_inline(text: str) -> list[tuple[str, str, str | None]]:
    """-> list of (kind, text, url) where kind is plain|bold|italic|link."""
    spans: list[tuple[str, str, str | None]] = []
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            spans.append(("plain", text[pos : m.start()], None))
        if m.group(1) is not None:
            spans.append(("bold", m.group(1), None))
        elif m.group(2) is not None:
            spans.append(("italic", m.group(2), None))
        else:
            spans.append(("link", m.group(3), m.group(4)))
        pos = m.end()
    if pos < len(text):
        spans.append(("plain", text[pos:], None))
    return spans


def parse_blocks(markdown: str) -> list[tuple[str, str]]:
    """-> list of (kind, text) where kind is h1|h2|h3|bullet|numbered|para.
    For numbered items the text keeps its "N. " prefix."""
    blocks: list[tuple[str, str]] = []
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            blocks.append(("h3", line[4:]))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:]))
        elif line.startswith("# "):
            blocks.append(("h1", line[2:]))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:]))
        elif _NUMBERED.match(line):
            blocks.append(("numbered", line))
        else:
            blocks.append(("para", line))
    return blocks


# ----------------------------- DOCX -----------------------------------------


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _fill_docx_paragraph(paragraph, text: str) -> None:
    for kind, value, url in parse_inline(text):
        if kind == "link" and url:
            _add_hyperlink(paragraph, value, url)
        else:
            run = paragraph.add_run(value)
            run.bold = kind == "bold"
            run.italic = kind == "italic"


def render_docx(title: str, markdown: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.core_properties.title = title
    for kind, text in parse_blocks(markdown):
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(text, level=int(kind[1]))
        elif kind == "bullet":
            _fill_docx_paragraph(doc.add_paragraph(style="List Bullet"), text)
        else:
            # "numbered" lines keep their own "N. " prefix as plain text, so
            # numbering restarts correctly in every list (Word's automatic
            # numbering would continue counting across lists).
            _fill_docx_paragraph(doc.add_paragraph(), text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------ PDF -----------------------------------------


def _latin(text: str) -> str:
    return text.encode("cp1252", "replace").decode("cp1252")


def _pdf_markup(text: str) -> str:
    out = []
    for kind, value, url in parse_inline(_latin(text)):
        value = escape(value)
        if kind == "bold":
            out.append(f"<b>{value}</b>")
        elif kind == "italic":
            out.append(f"<i>{value}</i>")
        elif kind == "link" and url:
            out.append(f'<a href="{escape(url, {chr(34): "&quot;"})}" color="#0563C1"><u>{value}</u></a>')
        else:
            out.append(value)
    return "".join(out)


def render_pdf(title: str, markdown: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10, leading=14, spaceAfter=4)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=14, bulletIndent=4)
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=20, leading=24, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=14, leading=18, spaceBefore=10, spaceAfter=4)
    h3 = ParagraphStyle("h3", parent=styles["Heading3"], fontSize=11, leading=14, spaceBefore=6, spaceAfter=2)
    numbered = ParagraphStyle("numbered", parent=body, leftIndent=14, firstLineIndent=-14)

    story = []
    for kind, text in parse_blocks(markdown):
        if kind == "h1":
            story.append(Paragraph(_pdf_markup(text), h1))
        elif kind == "h2":
            story.append(Paragraph(_pdf_markup(text), h2))
        elif kind == "h3":
            story.append(Paragraph(_pdf_markup(text), h3))
        elif kind == "bullet":
            story.append(Paragraph(_pdf_markup(text), bullet, bulletText="•"))
        elif kind == "numbered":
            story.append(Paragraph(_pdf_markup(text), numbered))
        else:
            story.append(Paragraph(_pdf_markup(text), body))
    story.append(Spacer(1, 6 * mm))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(18 * mm, 10 * mm, "CISE - Crisis Intelligence Synthesis Engine")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
        title=_latin(title), author="CISE",
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
